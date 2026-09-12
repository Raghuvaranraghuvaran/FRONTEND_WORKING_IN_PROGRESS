import os
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_complete_demo_docx():
    doc = Document()

    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("ReturnGuard • 7-Minute Demonstration Script: Exact Section Hierarchy")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)

        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("ReturnGuard • Line-Wise Section Guide: What to Show & What to Tell")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

    COLOR_NAVY = RGBColor(30, 27, 75)       # #1E1B4B
    COLOR_INDIGO = RGBColor(67, 56, 202)     # #4338CA
    COLOR_TEAL = RGBColor(13, 148, 136)      # #0D9488
    COLOR_DARK = RGBColor(30, 41, 59)        # #1E293B
    COLOR_BODY = RGBColor(51, 65, 85)        # #334155

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("RETURNGUARD")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(3)
    run_sub = p_sub.add_run("7-Minute Demonstration Guide: Line-Wise Section Hierarchy")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_INDIGO

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(12)
    run_meta = p_meta.add_run("Exact Hierarchy Order for All Merchant and Shopper Sections\nFormatted with Exact Heading Names, Routes, Time Budgets, What to Show, What to Tell & Technical Notes")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    # Master Schedule Table
    p_sched = doc.add_paragraph()
    p_sched.paragraph_format.space_before = Pt(4)
    p_sched.paragraph_format.space_after = Pt(4)
    r_sched = p_sched.add_run("Hierarchical Master Schedule (All 32 Sections Line-Wise)")
    r_sched.font.name = "Calibri"
    r_sched.font.size = Pt(11)
    r_sched.font.bold = True
    r_sched.font.color.rgb = COLOR_NAVY

    schedule_data = [
        # MERCHANT SECTIONS (Total: 5:00)
        ("1", "Merchant Dashboard", "Merchant", "/merchant/dashboard", "0:00 – 0:15", "15s", "Store KPIs, Return Rate (18.5%), Risk Tier"),
        ("2", "Flagged Cases", "Merchant", "/merchant/flagged-cases", "0:15 – 1:00", "45s", "Priority P: 1000, Actual photos, +1 badge, Review"),
        ("3", "Flagged Case Investigation Detail", "Merchant", "/merchant/flagged-cases/:id", "1:00 – 1:25", "25s", "Forensic dossier, serial mismatch, 1-click decisions"),
        ("4", "Delivery Agents", "Merchant", "/merchant/delivery-agents", "1:25 – 2:30", "65s", "Courier collusion, +26% Anomaly Gap, route audits"),
        ("5", "Fraud Rules & Config", "Merchant", "/merchant/fraud-config", "2:30 – 3:35", "65s", "Industry presets, sliders, VIP lists, live simulator"),
        ("6", "Customer Directory & Dossiers", "Merchant", "/merchant/customers", "3:35 – 3:50", "15s", "Behavioral dossier, progressive friction (Block COD)"),
        ("7", "Products Catalog", "Merchant", "/merchant/products", "3:50 – 4:00", "10s", "16 seeded items, Low Stock (≤5), Bulk CSV Upload"),
        ("8", "Orders Management", "Merchant", "/merchant/orders", "4:00 – 4:10", "10s", "Dynamic tabs (Delivered 6, COD 6), fulfillment actions"),
        ("9", "Coupons & Promotions", "Merchant", "/merchant/coupons", "4:10 – 4:20", "10s", "Active promos, live toggle switch via PATCH requests"),
        ("10", "Immutable Audit Log", "Merchant", "/merchant/audit-log", "4:20 – 4:30", "10s", "Enterprise non-repudiation: Actor, Action, Target, Notes"),
        ("11", "Setup & Multi-Tenancy Onboarding", "Merchant", "/merchant/onboarding", "4:30 – 4:38", "8s", "Assigned Merchant ID, store slug, isolated DB context"),
        ("12", "Analytics", "Merchant", "/merchant/analytics", "4:38 – 4:46", "8s", "Return rate trends, fraud recovery vs losses metrics"),
        ("13", "Merchant Settings", "Merchant", "/merchant/settings", "4:46 – 4:53", "7s", "Return window policies, notification thresholds"),
        ("14", "Merchant Login", "Merchant", "/merchant/login", "4:53 – 4:57", "4s", "Secure tenant session issuance, JWT authorization"),
        ("15", "Merchant Register", "Merchant", "/merchant/register", "4:57 – 5:00", "3s", "New store onboarding, automated catalog bootstrap"),
        # SHOPPER SECTIONS (Total: 2:00)
        ("16", "Landing Page", "Shopper", "/", "5:00 – 5:08", "8s", "Dual-engine hero banner, fast access portals"),
        ("17", "Shopper Storefront", "Shopper", "/shop", "5:08 – 5:18", "10s", "Multi-category catalog, dynamic pricing, return badge"),
        ("18", "Product Detail", "Shopper", "/products/:id", "5:18 – 5:26", "8s", "High-res images, return policy window, serial gating"),
        ("19", "Shopping Cart", "Shopper", "/cart", "5:26 – 5:33", "7s", "Subtotal computation, promo codes, checkout dispatch"),
        ("20", "Shopper Wishlist", "Shopper", "/wishlist", "5:33 – 5:38", "5s", "Saved items for later purchase, 1-click cart move"),
        ("21", "Checkout", "Shopper", "/checkout", "5:38 – 5:58", "20s", "Dynamic COD gating: blocked for risky profiles"),
        ("22", "Payment Success", "Shopper", "/payment/success", "5:58 – 6:02", "4s", "Transaction authorization, order dispatch creation"),
        ("23", "Payment Failure", "Shopper", "/payment/failure", "6:02 – 6:06", "4s", "Failure telemetry, graceful retry fallback"),
        ("24", "Shopper Dashboard", "Shopper", "/dashboard", "6:06 – 6:14", "8s", "Recent orders overview, return claim status badge"),
        ("25", "Shopper Orders", "Shopper", "/orders", "6:14 – 6:24", "10s", "Fulfillment timeline gating: returns only after delivery"),
        ("26", "Shopper Return Wizard", "Shopper", "/orders/:id/return", "6:24 – 6:44", "20s", "3 steps: reason, serial check vs dispatch, photo proof"),
        ("27", "Return Tracking", "Shopper", "/returns/:id/track", "6:44 – 6:50", "6s", "Real-time claim status: review, pickup, inspection, refund"),
        ("28", "Shopper Notifications", "Shopper", "/notifications", "6:50 – 6:53", "3s", "Automated alert feeds: dispatch, approval, refund"),
        ("29", "Shopper Profile", "Shopper", "/profile", "6:53 – 6:56", "3s", "Account info, shipping addresses, security credentials"),
        ("30", "Help Center", "Shopper", "/help", "6:56 – 6:58", "2s", "Self-service return policy & doorstep unboxing guidelines"),
        ("31", "Shopper Login", "Shopper", "/login", "6:58 – 6:59", "1s", "Customer token authentication, order state sync"),
        ("32", "Shopper Register", "Shopper", "/register", "6:59 – 7:00", "1s", "New buyer profile creation with clean baseline score"),
    ]

    tbl_s = doc.add_table(rows=len(schedule_data) + 1, cols=6)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s.autofit = False
    col_w = [Inches(0.4), Inches(2.0), Inches(0.8), Inches(1.6), Inches(1.0), Inches(1.2)]
    for r in tbl_s.rows:
        for ci, w in enumerate(col_w):
            r.cells[ci].width = w

    headers = ["#", "Exact Heading Name", "Side", "Route / URL", "Timeline", "Key Focus"]
    for ci, h in enumerate(headers):
        c = tbl_s.cell(0, ci)
        set_cell_background(c, "1E293B")
        set_cell_margins(c, top=50, bottom=50, left=60, right=60)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, row in enumerate(schedule_data, 1):
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for ci, val in enumerate([row[0], row[1], row[2], row[3], row[4] + " (" + row[5] + ")", row[6]]):
            c = tbl_s.cell(idx, ci)
            set_cell_background(c, bg)
            set_cell_margins(c, top=35, bottom=35, left=50, right=50)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(7.5)
            if ci == 0:
                r.font.bold = True
                r.font.color.rgb = COLOR_INDIGO
            elif ci == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_DARK
            elif ci == 2:
                r.font.bold = True
                r.font.color.rgb = COLOR_TEAL if val == "Merchant" else RGBColor(190, 24, 93)
            elif ci == 4:
                r.font.bold = True
                r.font.color.rgb = COLOR_INDIGO
            else:
                r.font.color.rgb = COLOR_BODY

    doc.add_page_break()

    # Reusable Component Builders
    def add_section_banner(num, name, side, route, time_budget, badge="OPERATIONS"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.rows[0].cells[0].width = Inches(7.0)
        c = tbl.cell(0, 0)
        bg_color = "312E81" if side == "Merchant" else "831843"
        set_cell_background(c, bg_color)
        set_cell_margins(c, top=70, bottom=70, left=120, right=120)
        
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        
        r_num = p.add_run(f"SECTION {num}: [{side.upper()} SIDE]  ")
        r_num.font.name = "Calibri"
        r_num.font.bold = True
        r_num.font.size = Pt(9.5)
        r_num.font.color.rgb = RGBColor(199, 210, 254) if side == "Merchant" else RGBColor(254, 205, 211)

        r_badge = p.add_run(f"★ {badge.upper()} ★\n")
        r_badge.font.name = "Calibri"
        r_badge.font.bold = True
        r_badge.font.size = Pt(8.5)
        r_badge.font.color.rgb = RGBColor(253, 224, 71)

        # EXACT NAME AS HEADING
        r_title = p.add_run(f"{name.upper()}\n")
        r_title.font.name = "Calibri"
        r_title.font.bold = True
        r_title.font.size = Pt(13)
        r_title.font.color.rgb = RGBColor(255, 255, 255)

        r_sub = p.add_run(f"Route to Open: {route}   •   Time Allocation: {time_budget}")
        r_sub.font.name = "Calibri"
        r_sub.font.size = Pt(8.5)
        r_sub.font.color.rgb = RGBColor(224, 231, 255) if side == "Merchant" else RGBColor(255, 228, 230)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_show_block(bullets):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.rows[0].cells[0].width = Inches(7.0)
        c = tbl.cell(0, 0)
        set_cell_background(c, "EFF6FF")
        set_cell_margins(c, top=60, bottom=60, left=110, right=110)
        
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("🖥️  WHAT TO SHOW:")
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(29, 78, 216)

        for b in bullets:
            p_b = c.add_paragraph()
            p_b.paragraph_format.space_after = Pt(1.5)
            p_b.paragraph_format.left_indent = Inches(0.15)
            rb = p_b.add_run(f"• {b}")
            rb.font.name = "Calibri"
            rb.font.size = Pt(8.5)
            rb.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_say_block(speech_paragraphs):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.rows[0].cells[0].width = Inches(7.0)
        c = tbl.cell(0, 0)
        set_cell_background(c, "F0FDF4")
        set_cell_margins(c, top=60, bottom=60, left=110, right=110)
        
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("🎙️  WHAT TO TELL:")
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(21, 128, 61)

        for sp in speech_paragraphs:
            p_s = c.add_paragraph()
            p_s.paragraph_format.space_after = Pt(2)
            rs = p_s.add_run(f"\"{sp.strip()}\"")
            rs.font.name = "Calibri"
            rs.font.size = Pt(8.5)
            rs.font.italic = True
            rs.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_tech_box(backend_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.rows[0].cells[0].width = Inches(7.0)
        c = tbl.cell(0, 0)
        set_cell_background(c, "F8FAFC")
        set_cell_margins(c, top=50, bottom=50, left=110, right=110)
        
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        r1 = p.add_run("⚙️  Technical & Database Under the Hood: ")
        r1.font.name = "Calibri"
        r1.font.bold = True
        r1.font.size = Pt(8)
        r1.font.color.rgb = COLOR_INDIGO
        r2 = p.add_run(backend_text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8)
        r2.font.color.rgb = COLOR_BODY

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # PART 1: MERCHANT SECTIONS (1 to 15) - EXACT HIERARCHICAL ORDER
    # =========================================================================

    # 1. Merchant Dashboard
    add_section_banner(1, "Merchant Dashboard", "Merchant", "http://localhost:5173/merchant/dashboard", "0:00 – 0:15 (15s)", "EXECUTIVE OVERVIEW")
    add_show_block([
        "Open Executive Dashboard at http://localhost:5173/merchant/dashboard.",
        "Point out the 4 KPI metric cards: Total Orders (16), Gross Revenue (₹54,200), Flagged Cases Count, and Store Return Rate (18.5%).",
        "Highlight the Store Risk Tier (Medium): show that exceeding 25% return rate triggers automated escalation to High Risk.",
        "Scroll down to preview the Recent Flagged Returns queue table.",
    ])
    add_say_block([
        "We begin on the Merchant Executive Dashboard. This command center provides an immediate rollup of the store's financial health and return risk.",
        "Notice our live KPIs: total revenue, order count, and our store return rate of 18.5%. If return claims spike past 25%, ReturnGuard automatically elevates the store's risk tier, warning managers before capital reserves are impacted."
    ])
    add_tech_box("`GET /api/admin/dashboard/` executes aggregate SQL queries (`Sum('total')`, `Count('id')`) dynamically partitioned by tenant.")

    # 2. Flagged Cases
    add_section_banner(2, "Flagged Cases", "Merchant", "http://localhost:5173/merchant/flagged-cases", "0:15 – 1:00 (45s)", "CORE FRAUD PILLAR #1")
    add_show_block([
        "Navigate to Flagged Cases: http://localhost:5173/merchant/flagged-cases.",
        "Top 4 Metric Cards: Total Flagged Cases, Pending Review, Confirmed Fraud, Legitimate Claims.",
        "Filter Bar: Status pills (All Cases, Manual Review, Approved, Rejected), Risk Tier dropdown, and Cards vs Table view toggle.",
        "Inspect Case Cards: Financial Priority badge (P: 1000), customer name, order number, return reason, payment mode, and refund value.",
        "Product Thumbnails: Real product photos fetched directly from order items.",
        "Multi-Item Clarity: Clean '+1 more' or '+2 more' badge, with no redundant '+0' on single-item orders.",
        "Show the clean 'Review Case ->' action button.",
    ])
    add_say_block([
        "Pillar 1 is our Flagged Cases Command Center, prioritizing suspicious claims for review.",
        "Unlike standard systems that sort by submission date, ReturnGuard prioritizes by financial risk. Claims with high refund values or hardware serial mismatches receive a priority score of P: 1000, floating them straight to the top of the queue.",
        "Notice each card displays actual product photos from the order, clear return reasons, and multi-item indicators like '+1 more'. Single-item orders remain completely clean without clutter. Store admins immediately see the full scope of risk before taking action."
    ])
    add_tech_box("`GET /api/admin/returns/` serializes `ReturnRequest` and joins `Order.items` with product image URLs.")

    # 3. Flagged Case Investigation Detail
    add_section_banner(3, "Flagged Case Investigation Detail", "Merchant", "http://localhost:5173/merchant/flagged-cases/:id", "1:00 – 1:25 (25s)", "CORE FRAUD PILLAR #1 DETAIL")
    add_show_block([
        "Click 'Review Case ->' on a high-risk case to open the Forensic Investigation Modal.",
        "Show Risk Breakdown: Serial Mismatch (+40 pts), High Return Velocity (+30 pts).",
        "Inspect Customer History: 3 recent returns, 2 COD refusals.",
        "Show Evidence Photo Proof: Shopper-uploaded picture of packaging.",
        "Demonstrate 1-Click Action Buttons: Approve Return, Reject Return (Fraud Confirmed), Restrict Customer.",
    ])
    add_say_block([
        "Opening the case dossier reveals complete forensic evidence: exact risk signal breakdowns, previous customer return velocity, and customer-uploaded packaging photos.",
        "With one click, the merchant can Approve, Reject for fraud, or Restrict the customer. Rejecting atomically updates the return claim, marks the order, logs a confirmed violation on the shopper profile, and writes to our immutable audit log."
    ])
    add_tech_box("`POST /api/admin/returns/:id/review/` runs an atomic transaction across `ReturnRequest`, `Order`, `ShopperProfile`, and `AuditLog`.")

    # 4. Delivery Agents
    add_section_banner(4, "Delivery Agents", "Merchant", "http://localhost:5173/merchant/delivery-agents", "1:25 – 2:30 (65s)", "CORE FRAUD PILLAR #2")
    add_show_block([
        "Navigate to Delivery Agents: http://localhost:5173/merchant/delivery-agents.",
        "Top Summary Cards: Agents Monitored, High Risk / Under Investigation, Regional Baseline Return Rate, and Route Anomalies.",
        "Delivery Agents Table: Courier names (Suresh Kumar, Imran Khan, Pooja Nair), delivery hubs, and total route assignments.",
        "Highlight Anomaly Gap: Show agent Imran Khan with an Actual Return Rate of 38% against an expected regional baseline of 12% — a massive +26% Anomaly Gap!",
        "Show Risk Badges: Review, Monitor, Normal.",
        "Demonstrate Managerial Action: Click 'Investigate' to trigger a route audit, or 'Sign-off' to confirm physical inspection.",
    ])
    add_say_block([
        "Pillar 2 tackles one of e-commerce's largest blind spots: Transit Collusion and Rogue Courier Shrinkage.",
        "Often, missing accessories, package swapping, or 'customer refused' claims aren't caused by the buyer — they are perpetrated in transit by delivery personnel.",
        "ReturnGuard establishes a statistical baseline: the Expected Return Rate for every delivery hub. When a courier's actual return rate spikes abnormally above that baseline, ReturnGuard flags an Anomaly Gap.",
        "Look at agent Imran Khan: his return rate is 38% compared to the regional baseline of 12% — a +26% anomaly! Merchants can immediately flag the courier for formal route audits or log managerial inspections, protecting the business from internal shrinkage."
    ])
    add_tech_box("`DeliveryAgent` model calculates `anomaly_gap = return_rate - expected_return_rate`. Sign-offs recorded in `AgentRiskSnapshot`.")

    # 5. Fraud Rules & Config
    add_section_banner(5, "Fraud Rules & Config", "Merchant", "http://localhost:5173/merchant/fraud-config", "2:30 – 3:35 (65s)", "CORE FRAUD PILLAR #3")
    add_show_block([
        "Navigate to Fraud Rules: http://localhost:5173/merchant/fraud-config.",
        "Industry Presets: Click 'Electronics' (serial mismatch weight jumps to 40), then click 'Fashion' (wardrobing and velocity weights jump).",
        "Interactive Sliders: Drag sliders for Serial Mismatch, Return Frequency, and COD Refusals — show that slider track fills update in real time with live values.",
        "Risk Thresholds: Low (0–34), Medium (35–64), High (65–100).",
        "Switch to VIP & Rules List tab: Show Whitelist (VIP trusted buyers bypass evidence checks) and Blacklist (repeat offenders blocked).",
        "Live Risk Simulator (Right Sidebar): Enter ₹8,999 order value, toggle high return velocity, and show the live risk score recalculate dynamically.",
        "Click 'Save Changes' to demonstrate instant database persistence.",
    ])
    add_say_block([
        "Pillar 3 is our customizable Fraud Rules Engine, giving merchants complete control over risk thresholds.",
        "Merchants can choose pre-configured Industry Presets: Electronics prioritizes serial discrepancies, while Fashion focuses on wardrobing. Every single weight slider is interactive with live track fills.",
        "In the VIP Rules tab, merchants maintain Whitelists for VIP zero-friction returns and Blacklists to intercept known abusers. On the right, our Live Risk Simulator allows risk officers to test scenario parameters before deploying to production. Saving commits the rules immediately to the database."
    ])
    add_tech_box("`FraudConfiguration` stores JSON weights; `PATCH /api/admin/fraud-config/` updates weights and triggers audit logging.")

    # 6. Customer Directory & Dossiers
    add_section_banner(6, "Customer Directory & Dossiers", "Merchant", "http://localhost:5173/merchant/customers", "3:35 – 3:50 (15s)", "OPERATIONS")
    add_show_block([
        "Navigate to Customers: http://localhost:5173/merchant/customers.",
        "Customer Table: Name, Email, Phone, Orders, Returns, and Risk Tier (Low, Medium, High).",
        "Open high-risk customer Rohit Verma: show return rate, COD refusal frequency, and serial mismatch counts.",
        "Show Progressive Friction Actions: Restrict COD (Prepaid Only), Require Open-Box Unboxing, or Suspend Account.",
    ])
    add_say_block([
        "The Customers directory maintains behavioral profiles. Rather than outright banning suspicious buyers, ReturnGuard implements Progressive Friction: we mandate prepaid checkout or require doorstep unboxing proof, mitigating fraud while preserving legitimate sales."
    ])
    add_tech_box("`ShopperProfile` tracks behavioral statistics; `POST /api/fraud/customers/:id/action/` updates customer restrictions.")

    # 7. Products Catalog
    add_section_banner(7, "Products Catalog", "Merchant", "http://localhost:5173/merchant/products", "3:50 – 4:00 (10s)", "OPERATIONS")
    add_show_block([
        "Navigate to Products: http://localhost:5173/merchant/products.",
        "Show 16 seeded products spanning Ethnic Wear, Electronics, Footwear, and Decor.",
        "Click filter tabs: All Items (16), Low Stock (≤5) (shows 2 items), and Out of Stock (shows 1 item).",
        "Point out the Bulk CSV Upload tool for migrating store inventories.",
    ])
    add_say_block([
        "Products provides full catalog telemetry. Merchants can monitor stock levels, filter low-stock or out-of-stock items, and bulk upload product inventories via CSV."
    ])
    add_tech_box("`GET /api/admin/products/` with stock status queries; auto-seeds sample catalog for all merchant tenants.")

    # 8. Orders Management
    add_section_banner(8, "Orders Management", "Merchant", "http://localhost:5173/merchant/orders", "4:00 – 4:10 (10s)", "OPERATIONS")
    add_show_block([
        "Navigate to Orders: http://localhost:5173/merchant/orders.",
        "Show dynamic counter tabs: All (16), Delivered (6), In Transit (4), COD (6), Flagged (5).",
        "Click 'Delivered' or 'COD' to show instant table filtering.",
        "Show fulfillment actions: 'Mark Delivered' or 'Log Doorstep Refusal'.",
    ])
    add_say_block([
        "Orders handles fulfillment workflows. Marking an order 'Delivered' starts the customer return window, while logging doorstep refusals increments customer risk scores automatically."
    ])
    add_tech_box("`GET /api/admin/orders/` with item prefetches; `POST /api/admin/orders/:id/status/` updates fulfillment states.")

    # 9. Coupons & Promotions
    add_section_banner(9, "Coupons & Promotions", "Merchant", "http://localhost:5173/merchant/coupons", "4:10 – 4:20 (10s)", "OPERATIONS")
    add_show_block([
        "Navigate to Coupons: http://localhost:5173/merchant/coupons.",
        "Show coupon list: WELCOME10, FESTIVE20, FREESHIP.",
        "Toggle active status switch: show live switch toggle and success toast notification.",
    ])
    add_say_block([
        "Coupons manages promotional rules. The live status toggle triggers an instantaneous PATCH request, synchronizing discount validity across shopper checkouts in real time."
    ])
    add_tech_box("`GET /api/admin/coupons/` and `PATCH /api/admin/coupons/:id/` persist active promotional states.")

    # 10. Immutable Audit Log
    add_section_banner(10, "Immutable Audit Log", "Merchant", "http://localhost:5173/merchant/audit-log", "4:20 – 4:30 (10s)", "COMPLIANCE")
    add_show_block([
        "Navigate to Audit Log: http://localhost:5173/merchant/audit-log.",
        "Show audit table columns: Actor Email, Action, Target Entity, Context Notes, Timestamp.",
        "Point out recent events generated during our demo (return reviewed, coupon toggled).",
    ])
    add_say_block([
        "Every single administrative action is recorded in our Immutable Audit Log, establishing enterprise non-repudiation and SOC-2 / ISO compliance."
    ])
    add_tech_box("`AuditLog` records append-only entries through the centralized `log_action()` service.")

    # 11. Setup & Multi-Tenancy Onboarding
    add_section_banner(11, "Setup & Multi-Tenancy Onboarding", "Merchant", "http://localhost:5173/merchant/onboarding", "4:30 – 4:38 (8s)", "MULTI-TENANCY")
    add_show_block([
        "Navigate to Setup & API: http://localhost:5173/merchant/onboarding.",
        "Show assigned Merchant ID (RG-DEMO-0154AA), store slug, and webhook integration endpoints for Shopify / WooCommerce.",
    ])
    add_say_block([
        "ReturnGuard is a multi-tenant platform. Each merchant receives an isolated tenant partition identified by their unique Merchant ID, with webhook endpoints for third-party storefronts."
    ])
    add_tech_box("`Merchant` context enforced across all DRF views via `require_merchant_context()`.")

    # 12. Analytics
    add_section_banner(12, "Analytics", "Merchant", "http://localhost:5173/merchant/analytics", "4:38 – 4:46 (8s)", "ANALYTICS")
    add_show_block([
        "Navigate to Analytics: http://localhost:5173/merchant/analytics.",
        "Show Return Rate Trend Line, Fraud vs Legitimate Claim Resolution Ratios, and Capital Saved Metric.",
    ])
    add_say_block([
        "The Analytics dashboard tracks financial recovery, showing prevented fraud losses, customer dispute resolution ratios, and seasonal return trends."
    ])
    add_tech_box("Calculates rolling 30-day aggregate statistics from `ReturnRequest` and `Order` tables.")

    # 13. Merchant Settings
    add_section_banner(13, "Merchant Settings", "Merchant", "http://localhost:5173/merchant/settings", "4:46 – 4:53 (7s)", "SETTINGS")
    add_show_block([
        "Navigate to Settings: http://localhost:5173/merchant/settings.",
        "Show Store Details, Return Eligibility Window (e.g. 7 vs 14 days), and Notification Webhooks.",
    ])
    add_say_block([
        "Settings governs store-wide policies, such as default return eligibility windows and notification webhooks for merchant teams."
    ])
    add_tech_box("`MerchantSettings` persists store configuration per tenant.")

    # 14. Merchant Login
    add_section_banner(14, "Merchant Login", "Merchant", "http://localhost:5173/merchant/login", "4:53 – 4:57 (4s)", "AUTHENTICATION")
    add_show_block([
        "Show Merchant Login screen: http://localhost:5173/merchant/login.",
        "Show Merchant ID / Email credentials input and role-based login dispatch.",
    ])
    add_say_block([
        "Merchant Login verifies tenant credentials, issues a scoped JWT token, and binds the session to the store's private partition."
    ])
    add_tech_box("`POST /api/auth/merchant/login/` validates merchant membership and returns scoped JWT tokens.")

    # 15. Merchant Register
    add_section_banner(15, "Merchant Register", "Merchant", "http://localhost:5173/merchant/register", "4:57 – 5:00 (3s)", "AUTHENTICATION")
    add_show_block([
        "Show Merchant Registration: http://localhost:5173/merchant/register.",
        "Show Store Name, Business Email, and automated catalog bootstrap on onboarding.",
    ])
    add_say_block([
        "Merchant Registration bootstraps a new tenant workspace, auto-seeding baseline fraud rules, sample catalog items, and a unique Merchant ID."
    ])
    add_tech_box("`POST /api/auth/merchant/register/` creates `Merchant` record and triggers `ensure_merchant_sample_data()`.")

    doc.add_page_break()

    # =========================================================================
    # PART 2: SHOPPER SECTIONS (16 to 32) - EXACT HIERARCHICAL ORDER
    # =========================================================================

    # 16. Landing Page
    add_section_banner(16, "Landing Page", "Shopper", "http://localhost:5173/", "5:00 – 5:08 (8s)", "PUBLIC PLATFORM")
    add_show_block([
        "Open Landing Page at http://localhost:5173/.",
        "Show Hero Banner highlighting ReturnGuard's dual-engine protection for shoppers and merchants.",
        "Point out Quick Access navigation buttons to Storefront and Merchant Portal.",
    ])
    add_say_block([
        "Switching to the Shopper experience, our Landing Page introduces ReturnGuard's dual protection: fast, frictionless returns for honest shoppers, backed by automated fraud detection."
    ])
    add_tech_box("Public React marketing route featuring dynamic portal routing and platform value pillars.")

    # 17. Shopper Storefront
    add_section_banner(17, "Shopper Storefront", "Shopper", "http://localhost:5173/shop", "5:08 – 5:18 (10s)", "SHOPPING")
    add_show_block([
        "Navigate to Storefront: http://localhost:5173/shop.",
        "Browse product cards across Ethnic Wear, Electronics, and Footwear.",
        "Point out price tags, stock badges, and '7-Day Easy Return' guarantee pill.",
        "Click 'Add to Cart' on an item.",
    ])
    add_say_block([
        "On the customer storefront, shoppers browse products with live prices and clear return guarantees. Adding an item synchronizes state instantly to the cart."
    ])
    add_tech_box("`GET /api/products/` returns public active inventory with category and price filters.")

    # 18. Product Detail
    add_section_banner(18, "Product Detail", "Shopper", "http://localhost:5173/products/:productId", "5:18 – 5:26 (8s)", "SHOPPING")
    add_show_block([
        "Click on a product card to open Product Detail: http://localhost:5173/products/1.",
        "Show multi-image gallery, item specifications, stock availability, and return eligibility terms.",
    ])
    add_say_block([
        "The Product Detail view highlights item specifications, warranty details, and transparent return eligibility conditions."
    ])
    add_tech_box("`GET /api/products/:id/` fetches full product metadata and stock levels.")

    # 19. Shopping Cart
    add_section_banner(19, "Shopping Cart", "Shopper", "http://localhost:5173/cart", "5:26 – 5:33 (7s)", "SHOPPING")
    add_show_block([
        "Open Cart: http://localhost:5173/cart.",
        "Show item quantities, price subtotal, free shipping progress bar, and 'Proceed to Checkout' button.",
    ])
    add_say_block([
        "In the cart, customers review line items, apply coupon codes, and see real-time subtotal calculations before proceeding."
    ])
    add_tech_box("Client-side cart state synchronized with local storage and validated against live inventory.")

    # 20. Shopper Wishlist
    add_section_banner(20, "Shopper Wishlist", "Shopper", "http://localhost:5173/wishlist", "5:33 – 5:38 (5s)", "SHOPPING")
    add_show_block([
        "Navigate to Wishlist: http://localhost:5173/wishlist.",
        "Show saved favorite items with 1-click 'Move to Cart' functionality.",
    ])
    add_say_block([
        "The Wishlist allows shoppers to bookmark favorite products, moving them into the active cart with a single click."
    ])
    add_tech_box("`GET /api/wishlist/` persists saved user preferences.")

    # 21. Checkout
    add_section_banner(21, "Checkout", "Shopper", "http://localhost:5173/checkout", "5:38 – 5:58 (20s)", "PRE-PURCHASE FRAUD ENGINE")
    add_show_block([
        "Navigate to Checkout: http://localhost:5173/checkout.",
        "Show delivery address and order summary.",
        "Highlight Payment Methods: Point out UPI, Credit/Debit Card, and Cash on Delivery.",
        "Explain Risk Gating: Clean buyers can select COD freely; however, if the shopper profile has high COD refusals or excessive returns, COD is dynamically disabled with a 'Prepaid Only Required' badge.",
    ])
    add_say_block([
        "At checkout, ReturnGuard enforces Pre-Purchase Risk Gating.",
        "Clean shoppers enjoy seamless checkout with UPI, Card, or COD. But if our engine detects an account with repeated doorstep COD refusals or abusive return velocity, Cash on Delivery is automatically locked, requiring prepaid payment before order dispatch."
    ])
    add_tech_box("`POST /api/orders/` evaluates `ShopperProfile.risk_tier` and `total_cod_refusals` before authorizing COD.")

    # 22. Payment Success
    add_section_banner(22, "Payment Success", "Shopper", "http://localhost:5173/payment/success", "5:58 – 6:02 (4s)", "PAYMENT GATEWAY")
    add_show_block([
        "Show Payment Success screen: http://localhost:5173/payment/success.",
        "Show Transaction Reference, Order Confirmation number, and link to Orders Timeline.",
    ])
    add_say_block([
        "Upon successful payment authorization, the system displays transaction confirmation and provisions the order in the database."
    ])
    add_tech_box("Payment gateway webhook callback handles order confirmation and marks payment status as Paid.")

    # 23. Payment Failure
    add_section_banner(23, "Payment Failure", "Shopper", "http://localhost:5173/payment/failure", "6:02 – 6:06 (4s)", "PAYMENT GATEWAY")
    add_show_block([
        "Show Payment Failure screen: http://localhost:5173/payment/failure.",
        "Show error message, decline reason, and 'Retry Payment' button without losing cart items.",
    ])
    add_say_block([
        "If a transaction fails, ReturnGuard provides clear decline telemetry and allows shoppers to retry without cart abandonment."
    ])
    add_tech_box("Captures gateway error codes (insufficient funds, timeout) and logs payment attempt telemetry.")

    # 24. Shopper Dashboard
    add_section_banner(24, "Shopper Dashboard", "Shopper", "http://localhost:5173/dashboard", "6:06 – 6:14 (8s)", "CUSTOMER ACCOUNT")
    add_show_block([
        "Navigate to Shopper Dashboard: http://localhost:5173/dashboard.",
        "Show Recent Orders widget, Active Return Requests counter, and quick navigation links.",
    ])
    add_say_block([
        "The Shopper Dashboard gives customers a personalized hub summarizing their order status, active return claims, and account metrics."
    ])
    add_tech_box("`GET /api/shopper/dashboard/` aggregates order counts and active return statuses.")

    # 25. Shopper Orders
    add_section_banner(25, "Shopper Orders", "Shopper", "http://localhost:5173/orders", "6:14 – 6:24 (10s)", "CUSTOMER ACCOUNT")
    add_show_block([
        "Navigate to Orders: http://localhost:5173/orders.",
        "Show order cards with timeline steps: Order Placed -> Shipped -> Out for Delivery -> Delivered.",
        "Point out that 'Return Order' is only enabled on Delivered orders — in-transit orders cannot initiate fraudulent phantom returns.",
    ])
    add_say_block([
        "In Orders, shoppers follow live fulfillment tracking. Notice that Return Request buttons are strictly locked until delivery is confirmed by the courier, preventing phantom return fraud."
    ])
    add_tech_box("`GET /api/orders/` returns orders for authenticated shopper; enforces delivery eligibility checks.")

    # 26. Shopper Return Wizard
    add_section_banner(26, "Shopper Return Wizard", "Shopper", "http://localhost:5173/orders/:orderId/return", "6:24 – 6:44 (20s)", "POST-PURCHASE FRAUD ENGINE")
    add_show_block([
        "Click 'Return Order' on a delivered order to launch the 3-Step Wizard.",
        "Step 1: Select Return Reason (e.g. 'Defective / Damaged') and select specific item.",
        "Step 2: Serial / IMEI Input field — explain that ReturnGuard validates this number against warehouse dispatch records to block hardware swaps.",
        "Step 3: Photo Proof Upload — drag-and-drop zone requiring clear photo evidence of item and packaging condition.",
        "Click 'Submit Return Request' to generate live risk scoring.",
    ])
    add_say_block([
        "When submitting a return, ReturnGuard guides honest shoppers through a 3-step structured evidence flow.",
        "In Step 2, the shopper must enter the product's hardware serial or IMEI number. ReturnGuard verifies this against the warehouse dispatch record — if a customer attempts a serial swap, our AI engine flags it immediately.",
        "Step 3 collects photographic proof of packaging condition before routing the claim to the merchant's Flagged Cases queue."
    ])
    add_tech_box("`POST /api/returns/` validates `serial_number` against `OrderItem.serial_number` and handles multi-part photo uploads.")

    # 27. Return Tracking
    add_section_banner(27, "Return Tracking", "Shopper", "http://localhost:5173/returns/:returnId/track", "6:44 – 6:50 (6s)", "CUSTOMER ACCOUNT")
    add_show_block([
        "Show Return Tracking timeline: Claim Submitted -> Under Merchant Review -> Courier Pickup -> Warehouse Inspection -> Refund Issued.",
    ])
    add_say_block([
        "Customers track their return claim with transparent progress milestones from initial merchant review to warehouse inspection and refund disbursement."
    ])
    add_tech_box("`GET /api/returns/:id/track/` returns real-time claim status and inspection milestones.")

    # 28. Shopper Notifications
    add_section_banner(28, "Shopper Notifications", "Shopper", "http://localhost:5173/notifications", "6:50 – 6:53 (3s)", "CUSTOMER ACCOUNT")
    add_show_block([
        "Navigate to Notifications: http://localhost:5173/notifications.",
        "Show alert feeds for order confirmation, shipping updates, and return approvals.",
    ])
    add_say_block([
        "Shopper Notifications keeps buyers informed in real time when returns are approved and refunds are credited."
    ])
    add_tech_box("`GET /api/notifications/` fetches user-specific alert events.")

    # 29. Shopper Profile
    add_section_banner(29, "Shopper Profile", "Shopper", "http://localhost:5173/profile", "6:53 – 6:56 (3s)", "CUSTOMER ACCOUNT")
    add_show_block([
        "Navigate to Profile: http://localhost:5173/profile.",
        "Show customer contact details, saved delivery addresses, and account security settings.",
    ])
    add_say_block([
        "The Profile view allows shoppers to manage saved shipping addresses and update authentication credentials."
    ])
    add_tech_box("`GET /api/shopper/profile/` returns shopper profile and verified contact data.")

    # 30. Help Center
    add_section_banner(30, "Help Center", "Shopper", "http://localhost:5173/help", "6:56 – 6:58 (2s)", "PUBLIC PLATFORM")
    add_show_block([
        "Navigate to Help Center: http://localhost:5173/help.",
        "Show Return Policy FAQ, Doorstep Unboxing Guidelines, and Customer Support channels.",
    ])
    add_say_block([
        "The Help Center provides clear customer guidance on return eligibility, doorstep unboxing verification, and contact support."
    ])
    add_tech_box("Static policy rules and FAQ knowledge base.")

    # 31. Shopper Login
    add_section_banner(31, "Shopper Login", "Shopper", "http://localhost:5173/login", "6:58 – 6:59 (1s)", "AUTHENTICATION")
    add_show_block([
        "Show Shopper Login screen: http://localhost:5173/login.",
        "Show Email / Password authentication and JWT token handling.",
    ])
    add_say_block([
        "Shopper Login validates customer credentials and restores active cart and order history."
    ])
    add_tech_box("`POST /api/auth/login/` returns shopper JWT token.")

    # 32. Shopper Register
    add_section_banner(32, "Shopper Register", "Shopper", "http://localhost:5173/register", "6:59 – 7:00 (1s)", "AUTHENTICATION")
    add_show_block([
        "Show Shopper Registration screen: http://localhost:5173/register.",
        "Show new customer signup initializing a clean baseline risk score.",
    ])
    add_say_block([
        "Shopper Registration creates a verified customer account with a clean baseline score, concluding our demonstration."
    ])
    add_tech_box("`POST /api/auth/register/` creates `User` and `ShopperProfile` with default Low Risk status.")

    return doc

if __name__ == "__main__":
    doc = create_complete_demo_docx()
    out1 = "demo.docx"
    doc.save(out1)
    print("Saved:", os.path.abspath(out1))

    os.makedirs("Project Documents", exist_ok=True)
    try:
        out2 = os.path.join("Project Documents", "ReturnGuard_7min_Demo_Guide.docx")
        doc.save(out2)
        print("Saved:", os.path.abspath(out2))
    except Exception as e:
        print("Note on Project Documents/ReturnGuard_7min_Demo_Guide.docx lock:", e)

    try:
        out4 = os.path.join("Project Documents", "ReturnGuard_Demo_Hierarchical_Guide.docx")
        doc.save(out4)
        print("Saved:", os.path.abspath(out4))
    except Exception as e:
        print("Note on out4:", e)
