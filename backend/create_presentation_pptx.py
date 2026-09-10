import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Theme colors - Light Background Template
    c_light_bg = RGBColor(248, 250, 252)    # #F8FAFC (Clean soft canvas)
    c_card_white = RGBColor(255, 255, 255)  # #FFFFFF (Pure card surface)
    c_hero_bg = RGBColor(241, 245, 249)     # #F1F5F9 (Subtle off-white)
    c_indigo = RGBColor(67, 56, 202)        # #4338CA (Deep Indigo accent)
    c_teal = RGBColor(13, 148, 136)         # #0D9488 (Emerald Teal brand)
    c_text_dark = RGBColor(15, 23, 42)      # #0F172A (Deep Slate text)
    c_text_muted = RGBColor(71, 85, 105)    # #475569 (Secondary text)
    c_critical = RGBColor(220, 38, 38)      # #DC2626 (Alert Red)
    c_border = RGBColor(203, 213, 225)      # #CBD5E1 (Light Slate border)
    c_accent_stripe = RGBColor(13, 148, 136)# #0D9488 (Top template bar)

    base_dir = os.path.abspath("screenshots")
    img_landing = os.path.join(base_dir, "real_landing_page.png")
    img_shopper_dash = os.path.join(base_dir, "shopper_dashboard.png")
    img_dashboard = os.path.join(base_dir, "3_merchant_dashboard.png")
    img_flagged_queue = os.path.join(base_dir, "4_flagged_cases_queue.png")
    img_case_detail = os.path.join(base_dir, "5_flagged_case_detail_28checkpoints.png")
    img_tech_logos = os.path.join(base_dir, "tech_stack", "official_tech_logos.png")

    def apply_slide_template(slide):
        """Applies a consistent, polished light background template to every slide."""
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_light_bg
        bg.line.fill.background()

        # Top template accent bar (0.07" height)
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.07))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = c_accent_stripe
        top_bar.line.fill.background()

    def add_left_header(slide, title_text, slide_label):
        """Header with slide label on the left side (no 'of X')."""
        num_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(1.3), Inches(0.36))
        num_box.fill.solid()
        num_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
        num_box.line.color.rgb = c_border
        tf_num = num_box.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_label.upper()
        p_num.font.size = Pt(10.5)
        p_num.font.bold = True
        p_num.font.color.rgb = c_teal
        p_num.alignment = PP_ALIGN.CENTER

        title_box = slide.shapes.add_textbox(Inches(2.3), Inches(0.35), Inches(10.1), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = c_text_dark

    def add_highlight_bullet(tf, title, desc, space_after=8):
        p = tf.add_paragraph() if tf.paragraphs[0].text else tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = title + " "
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = c_text_dark

        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(9)
        r2.font.color.rgb = c_text_muted
        p.space_after = Pt(space_after)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (LIGHT BACKGROUND TEMPLATE)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide1)

    s1_num = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(1.3), Inches(0.36))
    s1_num.fill.solid()
    s1_num.fill.fore_color.rgb = RGBColor(241, 245, 249)
    s1_num.line.color.rgb = c_border
    tf_s1 = s1_num.text_frame
    p_s1 = tf_s1.paragraphs[0]
    p_s1.text = "SLIDE 1"
    p_s1.font.size = Pt(10.5)
    p_s1.font.bold = True
    p_s1.font.color.rgb = c_teal
    p_s1.alignment = PP_ALIGN.CENTER

    # Hero card on light canvas
    hero_card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.5), Inches(10.9), Inches(5.0))
    hero_card1.fill.solid()
    hero_card1.fill.fore_color.rgb = c_card_white
    hero_card1.line.color.rgb = c_border

    tf1 = hero_card1.text_frame
    tf1.margin_left = Inches(0.5)
    tf1.margin_right = Inches(0.5)
    tf1.margin_top = Inches(0.6)
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs[0]
    p1.text = "ReturnGuard"
    p1.font.size = Pt(52)
    p1.font.bold = True
    p1.font.color.rgb = c_text_dark
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(12)

    p2 = tf1.add_paragraph()
    p2.text = "An Intelligent, COD-Native Risk Defense System for Indian D2C Commerce"
    p2.font.size = Pt(20)
    p2.font.color.rgb = c_text_muted
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(32)

    p_meta = tf1.add_paragraph()
    r_team = p_meta.add_run()
    r_team.text = "Project Team: Team ReturnGuard       |       "
    r_team.font.size = Pt(15)
    r_team.font.bold = True
    r_team.font.color.rgb = c_teal

    r_mentor = p_meta.add_run()
    r_mentor.text = "Mentor Name: Thavayee"
    r_mentor.font.size = Pt(15)
    r_mentor.font.bold = True
    r_mentor.font.color.rgb = c_text_dark
    p_meta.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2: PROJECT OVERVIEW (REPLACED MERCHANT DASHBOARD WITH STOREFRONT LANDING PAGE)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide2)
    add_left_header(slide2, "PROJECT OVERVIEW", "Slide 2")

    left_box2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.7))
    left_box2.fill.solid()
    left_box2.fill.fore_color.rgb = c_card_white
    left_box2.line.color.rgb = c_border
    tf2 = left_box2.text_frame
    tf2.margin_left = Inches(0.28)
    tf2.margin_right = Inches(0.28)
    tf2.margin_top = Inches(0.28)
    tf2.word_wrap = True

    add_highlight_bullet(tf2, "• Problem Context in Indian D2C:", "Cash-on-Delivery (COD) represents 60%+ of Indian e-commerce orders, creating severe financial exposure where merchants pay 3x forward freight on illegitimate returns.", 10)
    add_highlight_bullet(tf2, "• Dual-Portal Ecosystem:", "A complete, production-ready web platform providing a customer storefront alongside an operations portal for merchant fraud investigation.", 10)
    add_highlight_bullet(tf2, "• Real-Time Risk Engine:", "Evaluates all incoming transactions and return claims dynamically across a 0 to 100 risk scoring spectrum in sub-second execution.", 10)
    add_highlight_bullet(tf2, "• Target Beneficiaries:", "Small-to-mid Indian D2C retailers ($50K–$5M ARR) selling apparel, lifestyle, and electronics who cannot afford costly enterprise fraud solutions.", 10)
    add_highlight_bullet(tf2, "• Frictionless by Default:", "85%+ of genuine shoppers enjoy instant zero-friction auto-approvals; verification escalates only when specific risk anomalies trigger.", 10)
    add_highlight_bullet(tf2, "• Margin Preservation Value:", "Converts reverse logistics from an uncontrolled profit drain into a secure, revenue-protecting operational asset.", 10)

    # Replaced with Storefront Landing Page
    if os.path.exists(img_landing):
        slide2.shapes.add_picture(img_landing, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.2))
        cap2 = slide2.shapes.add_textbox(Inches(6.8), Inches(6.55), Inches(5.7), Inches(0.4))
        tf_cap2 = cap2.text_frame
        p_cap2 = tf_cap2.paragraphs[0]
        p_cap2.text = "Live Screen: Storefront Landing Page & Shopper Commerce Experience"
        p_cap2.font.size = Pt(9.5)
        p_cap2.font.color.rgb = c_text_muted
        p_cap2.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 3: SHOPPER & MERCHANT PLATFORM DASHBOARDS
    # (Shows Shopper Dashboard + REAL Merchant Operations Dashboard)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide3)
    add_left_header(slide3, "PLATFORM ARCHITECTURE: SHOPPER & MERCHANT DASHBOARDS", "Slide 3")

    col_w = Inches(5.7)

    # Left: Shopper Dashboard
    if os.path.exists(img_shopper_dash):
        slide3.shapes.add_picture(img_shopper_dash, Inches(0.8), Inches(1.3), col_w, Inches(2.7))
        cap_shop = slide3.shapes.add_textbox(Inches(0.8), Inches(4.02), col_w, Inches(0.35))
        tf_cs = cap_shop.text_frame
        p_cs = tf_cs.paragraphs[0]
        p_cs.text = "Shopper Portal: Order History, Reward Points & 1-Click Return/Exchange Hub"
        p_cs.font.size = Pt(9)
        p_cs.font.color.rgb = c_text_muted
        p_cs.alignment = PP_ALIGN.CENTER

    box_shop = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.4), col_w, Inches(2.65))
    box_shop.fill.solid()
    box_shop.fill.fore_color.rgb = c_card_white
    box_shop.line.color.rgb = c_border
    tf_bs = box_shop.text_frame
    tf_bs.margin_left = Inches(0.25)
    tf_bs.margin_right = Inches(0.25)
    tf_bs.margin_top = Inches(0.2)
    tf_bs.word_wrap = True

    p_bsh = tf_bs.paragraphs[0]
    p_bsh.text = "SHOPPER STOREFRONT & SELF-SERVICE PORTAL"
    p_bsh.font.size = Pt(11.5)
    p_bsh.font.bold = True
    p_bsh.font.color.rgb = c_indigo
    p_bsh.space_after = Pt(6)

    add_highlight_bullet(tf_bs, "• 1-Click Return Initiation:", "Seamless self-service return and exchange flow directly from customer order history with dynamic reason collection.", 6)
    add_highlight_bullet(tf_bs, "• Smart Exchange Incentive:", "Prompts instant size/color exchanges over cash refunds, reducing reverse freight fees while preserving the sale.", 6)
    add_highlight_bullet(tf_bs, "• Live Order & Return Tracking:", "Transparent real-time status tracking showing pickup scheduling, doorstep verification, and refund timeline.", 6)
    add_highlight_bullet(tf_bs, "• Friction-Free Experience:", "85%+ of genuine shoppers enjoy instant auto-approval with zero administrative hurdles or hold times.", 6)

    # Right: Real Merchant Dashboard
    if os.path.exists(img_dashboard):
        slide3.shapes.add_picture(img_dashboard, Inches(6.8), Inches(1.3), col_w, Inches(2.7))
        cap_merch = slide3.shapes.add_textbox(Inches(6.8), Inches(4.02), col_w, Inches(0.35))
        tf_cm = cap_merch.text_frame
        p_cm = tf_cm.paragraphs[0]
        p_cm.text = "Merchant Portal: Executive Risk KPIs, Flagged Returns Queue & Fraud Triage"
        p_cm.font.size = Pt(9)
        p_cm.font.color.rgb = c_text_muted
        p_cm.alignment = PP_ALIGN.CENTER

    box_merch = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.4), col_w, Inches(2.65))
    box_merch.fill.solid()
    box_merch.fill.fore_color.rgb = c_card_white
    box_merch.line.color.rgb = c_border
    tf_bm = box_merch.text_frame
    tf_bm.margin_left = Inches(0.25)
    tf_bm.margin_right = Inches(0.25)
    tf_bm.margin_top = Inches(0.2)
    tf_bm.word_wrap = True

    p_bmh = tf_bm.paragraphs[0]
    p_bmh.text = "MERCHANT OPERATIONS COMMAND CENTER"
    p_bmh.font.size = Pt(11.5)
    p_bmh.font.bold = True
    p_bmh.font.color.rgb = c_teal
    p_bmh.space_after = Pt(6)

    add_highlight_bullet(tf_bm, "• Executive KPI Intelligence:", "Real-time visibility into overall return rates, prevented loss values, risk tier breakdown, and order velocity.", 6)
    add_highlight_bullet(tf_bm, "• Automated Anomaly Triage:", "Continuous 0–100 risk scoring automatically funneling suspicious claims into a prioritized review queue.", 6)
    add_highlight_bullet(tf_bm, "• Configurable Fraud Policies:", "Dynamic sliders to tune sensitivity per category, seasonal festive wardrobing rules, and serial tracking.", 6)
    add_highlight_bullet(tf_bm, "• Custody Chain & Logistics:", "Integrates doorstep digital witnessing, courier handover verification, and driver collusion tracking.", 6)

    # =========================================================================
    # SLIDE 4: CHALLENGES & PRACTICAL SOLUTIONS
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide4)
    add_left_header(slide4, "CHALLENGES & SOLUTIONS", "Slide 4")

    left_box4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.7))
    left_box4.fill.solid()
    left_box4.fill.fore_color.rgb = c_card_white
    left_box4.line.color.rgb = c_border
    tf4 = left_box4.text_frame
    tf4.margin_left = Inches(0.28)
    tf4.margin_right = Inches(0.28)
    tf4.margin_top = Inches(0.28)
    tf4.word_wrap = True

    add_highlight_bullet(tf4, "• Challenge 1 — Wardrobing & Festive Rental Abuse:", "Shoppers buy premium ethnic apparel for Diwali or wedding seasons, wear them once, and return for full refunds. Solution (CP4): Seasonal cultural risk tightening, mandatory tag-intact verification, and pre-return condition checks.", 11)
    add_highlight_bullet(tf4, "• Challenge 2 — Size Bracketing Multi-Orders:", "Buyers purchase 3–4 sizes of the same SKU intending to keep only one, inflating freight costs by 300%. Solution (CP10): Cart-level multi-size detection with instant size exchange incentives over cash returns.", 11)
    add_highlight_bullet(tf4, "• Challenge 3 — Electronics Hardware Swapping:", "High-value electronics returned with missing chargers, clone units, or dummy ballast weights. Solution (CP17): Strict outbound vs inbound serial/IMEI verification with automated +50 pt critical risk flagging.", 11)
    add_highlight_bullet(tf4, "• Challenge 4 — Courier & Driver Collusion:", "Delivery drivers falsely marking items as 'customer rejected' or siphoning returned goods during transit. Solution (CP24): Geo-stamped doorstep signatures, tamper-evident seals, and driver route anomaly analytics.", 11)

    if os.path.exists(img_flagged_queue):
        slide4.shapes.add_picture(img_flagged_queue, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.2))
        cap4 = slide4.shapes.add_textbox(Inches(6.8), Inches(6.55), Inches(5.7), Inches(0.4))
        tf_cap4 = cap4.text_frame
        p_cap4 = tf_cap4.paragraphs[0]
        p_cap4.text = "Live Screen: Flagged Cases Queue (Actionable Risk Scoring & Resolution Workflow)"
        p_cap4.font.size = Pt(9.5)
        p_cap4.font.color.rgb = c_text_muted
        p_cap4.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 5: KEY FEATURES
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide5)
    add_left_header(slide5, "KEY FEATURES", "Slide 5")

    left_box5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.7))
    left_box5.fill.solid()
    left_box5.fill.fore_color.rgb = c_card_white
    left_box5.line.color.rgb = c_border
    tf5 = left_box5.text_frame
    tf5.margin_left = Inches(0.28)
    tf5.margin_right = Inches(0.28)
    tf5.margin_top = Inches(0.28)
    tf5.word_wrap = True

    add_highlight_bullet(tf5, "• Customer & Risk Profile Context:", "Real-time shopper telemetry displaying customer identity, order history, active return rate, and flagged risk tier (High Risk • 78 pts).", 10)
    add_highlight_bullet(tf5, "• 5-Step Return & Risk Engine:", "Automated lifecycle progression through Return Requested → Pickup Scheduled → Product Inspection → Risk Evaluation → Merchant Decision.", 10)
    add_highlight_bullet(tf5, "• 1-Click Merchant Decision Panel:", "Direct operational execution buttons: Approve Refund (Green), Store Credit +5% (Blue), Replacement (Indigo), Hold (Amber), Reject (Rose), or Reject + Restrict (Dark Red).", 10)
    add_highlight_bullet(tf5, "• 4-Tier / 28 Risk Checkpoints:", "Comprehensive evaluation pipeline covering Pre-Scoring Gates, Physical Attributes, Behavioral Biometrics, and Composite Decision Rules.", 10)
    add_highlight_bullet(tf5, "• 6-Level Escalation Ladder:", "Proportionate customer friction scaling from instant auto-approval to OTP challenge, COD lock, and account termination.", 10)
    add_highlight_bullet(tf5, "• Full Audit & Case Documentation:", "Structured decision notes capture and immutable audit logging for fraud disputes and chargeback defense.", 10)

    if os.path.exists(img_case_detail):
        slide5.shapes.add_picture(img_case_detail, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.2))
        cap5 = slide5.shapes.add_textbox(Inches(6.8), Inches(6.55), Inches(5.7), Inches(0.4))
        tf_cap5 = cap5.text_frame
        p_cap5 = tf_cap5.paragraphs[0]
        p_cap5.text = "Live Screen: Flagged Case Review (Customer Profile, 5-Step Return Engine & Decision Panel)"
        p_cap5.font.size = Pt(9.5)
        p_cap5.font.color.rgb = c_text_muted
        p_cap5.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 6: TECHNOLOGY STACK
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide6)
    add_left_header(slide6, "TECHNOLOGY STACK", "Slide 6")

    if os.path.exists(img_tech_logos):
        slide6.shapes.add_picture(img_tech_logos, Inches(2.2), Inches(1.25), Inches(8.93), Inches(1.9))

    bot_left6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.3), Inches(5.7), Inches(3.7))
    bot_left6.fill.solid()
    bot_left6.fill.fore_color.rgb = c_card_white
    bot_left6.line.color.rgb = c_border
    tf6_l = bot_left6.text_frame
    tf6_l.margin_left = Inches(0.28)
    tf6_l.margin_right = Inches(0.28)
    tf6_l.margin_top = Inches(0.25)
    tf6_l.word_wrap = True

    p6_lh = tf6_l.paragraphs[0]
    p6_lh.text = "FRONTEND & CLIENT ARCHITECTURE"
    p6_lh.font.size = Pt(13)
    p6_lh.font.bold = True
    p6_lh.font.color.rgb = c_text_dark
    p6_lh.space_after = Pt(8)

    add_highlight_bullet(tf6_l, "• React & Vite SPA:", "Modular component architecture powering both shopper store and merchant command center with hot module replacement.", 8)
    add_highlight_bullet(tf6_l, "• Tailwind CSS Design System:", "Utility-first design delivering responsive layouts, risk badges, and accessible data tables.", 8)
    add_highlight_bullet(tf6_l, "• Lucide React Icons & Context API:", "Clean icon set with lightweight global state management for auth, cart, and live alerts.", 8)
    add_highlight_bullet(tf6_l, "• Resilient API Client:", "Axios HTTP layer featuring automatic JWT token refreshing and error boundary fallback states.", 8)

    bot_right6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.3), Inches(5.7), Inches(3.7))
    bot_right6.fill.solid()
    bot_right6.fill.fore_color.rgb = c_card_white
    bot_right6.line.color.rgb = c_border
    tf6_r = bot_right6.text_frame
    tf6_r.margin_left = Inches(0.28)
    tf6_r.margin_right = Inches(0.28)
    tf6_r.margin_top = Inches(0.25)
    tf6_r.word_wrap = True

    p6_rh = tf6_r.paragraphs[0]
    p6_rh.text = "BACKEND, DATABASE & TASK PIPELINE"
    p6_rh.font.size = Pt(13)
    p6_rh.font.bold = True
    p6_rh.font.color.rgb = c_text_dark
    p6_rh.space_after = Pt(8)

    add_highlight_bullet(tf6_r, "• Python & Django REST Framework:", "Robust RESTful API architecture running 28-checkpoint scoring with sub-second execution.", 8)
    add_highlight_bullet(tf6_r, "• PostgreSQL with Row-Level Security:", "Relational database with strict tenant isolation, foreign keys, and ACID compliance.", 8)
    add_highlight_bullet(tf6_r, "• Redis & Celery Worker Queue:", "Distributed asynchronous background pipeline handling risk scoring and automated alerts.", 8)
    add_highlight_bullet(tf6_r, "• Security & Invoicing Engine:", "HMAC SHA-256 webhook signatures, encrypted telemetry, and ReportLab PDF invoice generation.", 8)

    # =========================================================================
    # SLIDE 7: CONCLUSION
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide7)
    add_left_header(slide7, "CONCLUSION", "Slide 7")

    c_w = Inches(5.7)
    c_h = Inches(5.7)

    box7_l = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), c_w, c_h)
    box7_l.fill.solid()
    box7_l.fill.fore_color.rgb = c_card_white
    box7_l.line.color.rgb = c_border
    tf7_l = box7_l.text_frame
    tf7_l.margin_left = Inches(0.28)
    tf7_l.margin_right = Inches(0.28)
    tf7_l.margin_top = Inches(0.28)
    tf7_l.word_wrap = True

    p7_lh = tf7_l.paragraphs[0]
    p7_lh.text = "CORE SYSTEM DELIVERABLES"
    p7_lh.font.size = Pt(13)
    p7_lh.font.bold = True
    p7_lh.font.color.rgb = c_text_dark
    p7_lh.space_after = Pt(10)

    add_highlight_bullet(tf7_l, "• Complete Working Platform:", "Dual-portal e-commerce application integrating a customer storefront with a real-time merchant operations command portal.", 10)
    add_highlight_bullet(tf7_l, "• 28 Algorithmic Checkpoints:", "Comprehensive multi-tier scoring pipeline analyzing behavioral, cultural, physical, and logistics checks.", 10)
    add_highlight_bullet(tf7_l, "• 6-Level Escalation Ladder:", "Dynamic, proportional friction scaling from automated approval to doorstep OTP, COD restriction, and blacklisting.", 10)
    add_highlight_bullet(tf7_l, "• Production API Architecture:", "Scalable Django REST backend with asynchronous task workers and PostgreSQL Row-Level Security isolation.", 10)
    add_highlight_bullet(tf7_l, "• Full Audit & Case History:", "Immutable tracking logs capturing every risk signal, customer action, and merchant resolution.", 10)

    box7_r = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), c_w, c_h)
    box7_r.fill.solid()
    box7_r.fill.fore_color.rgb = c_card_white
    box7_r.line.color.rgb = c_border
    tf7_r = box7_r.text_frame
    tf7_r.margin_left = Inches(0.28)
    tf7_r.margin_right = Inches(0.28)
    tf7_r.margin_top = Inches(0.28)
    tf7_r.word_wrap = True

    p7_rh = tf7_r.paragraphs[0]
    p7_rh.text = "BUSINESS VALUE & QUANTIFIED IMPACT"
    p7_rh.font.size = Pt(13)
    p7_rh.font.bold = True
    p7_rh.font.color.rgb = c_teal
    p7_rh.space_after = Pt(10)

    add_highlight_bullet(tf7_r, "• 85%+ Instant Approvals:", "Genuine customers experience zero friction with instant auto-approval on legitimate orders and return claims.", 10)
    add_highlight_bullet(tf7_r, "• 70% Cut in Wardrobing Abuse:", "Cultural and seasonal calendar rules eliminate organized festive garment rental fraud during weddings and festivals.", 10)
    add_highlight_bullet(tf7_r, "• Reverse Freight Recovery:", "Eliminates unnecessary 3x forward logistics fees on abusive multi-size orders and merchandise swapping.", 10)
    add_highlight_bullet(tf7_r, "• Protects Merchant Margin:", "Empowers small-to-mid Indian D2C retailers to offer safe COD payments without risking inventory bankruptcy.", 10)
    add_highlight_bullet(tf7_r, "• Granular Policy Customization:", "Store operators can adjust risk sensitivity sliders dynamically per product category to optimize margins.", 10)

    # =========================================================================
    # SLIDE 8: FUTURE SCOPE
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide8)
    add_left_header(slide8, "FUTURE SCOPE", "Slide 8")

    card8_w = Inches(5.7)
    card8_h = Inches(2.7)

    b8_1 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), card8_w, card8_h)
    b8_1.fill.solid()
    b8_1.fill.fore_color.rgb = c_card_white
    b8_1.line.color.rgb = c_border
    tf8_1 = b8_1.text_frame
    tf8_1.margin_left = Inches(0.25)
    tf8_1.margin_right = Inches(0.25)
    tf8_1.margin_top = Inches(0.22)
    tf8_1.word_wrap = True
    p8_1h = tf8_1.paragraphs[0]
    p8_1h.text = "1. MOBILE EDGE COMPUTER VISION"
    p8_1h.font.size = Pt(12)
    p8_1h.font.bold = True
    p8_1h.font.color.rgb = c_indigo
    p8_1h.space_after = Pt(6)
    add_highlight_bullet(tf8_1, "• Real-Time Doorstep Scanning:", "Deploy lightweight on-device AI vision models to couriers' handheld devices for real-time fabric texture analysis, tag integrity verification, and wear detection prior to acceptance.", 6)

    b8_2 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), card8_w, card8_h)
    b8_2.fill.solid()
    b8_2.fill.fore_color.rgb = c_card_white
    b8_2.line.color.rgb = c_border
    tf8_2 = b8_2.text_frame
    tf8_2.margin_left = Inches(0.25)
    tf8_2.margin_right = Inches(0.25)
    tf8_2.margin_top = Inches(0.22)
    tf8_2.word_wrap = True
    p8_2h = tf8_2.paragraphs[0]
    p8_2h.text = "2. CONSORTIUM FRAUD GRAPH NETWORK"
    p8_2h.font.size = Pt(12)
    p8_2h.font.bold = True
    p8_2h.font.color.rgb = c_teal
    p8_2h.space_after = Pt(6)
    add_highlight_bullet(tf8_2, "• Cross-Merchant Identity Mesh:", "Privacy-preserving zero-knowledge graph network that pools anonymized risk telemetry across independent Indian D2C stores to preemptively flag serial fraudsters.", 6)

    b8_3 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), card8_w, card8_h)
    b8_3.fill.solid()
    b8_3.fill.fore_color.rgb = c_card_white
    b8_3.line.color.rgb = c_border
    tf8_3 = b8_3.text_frame
    tf8_3.margin_left = Inches(0.25)
    tf8_3.margin_right = Inches(0.25)
    tf8_3.margin_top = Inches(0.22)
    tf8_3.word_wrap = True
    p8_3h = tf8_3.paragraphs[0]
    p8_3h.text = "3. MANIFEST V3 CHROME OPERATIONS EXTENSION"
    p8_3h.font.size = Pt(12)
    p8_3h.font.bold = True
    p8_3h.font.color.rgb = c_text_dark
    p8_3h.space_after = Pt(6)
    add_highlight_bullet(tf8_3, "• Embedded E-Commerce Integration:", "Browser toolbar extension injecting ReturnGuard's 0-100 risk badges and 1-click triage buttons directly inside Shopify, WooCommerce, and Magento order management panels.", 6)

    b8_4 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.3), card8_w, card8_h)
    b8_4.fill.solid()
    b8_4.fill.fore_color.rgb = c_card_white
    b8_4.line.color.rgb = c_border
    tf8_4 = b8_4.text_frame
    tf8_4.margin_left = Inches(0.25)
    tf8_4.margin_right = Inches(0.25)
    tf8_4.margin_top = Inches(0.22)
    tf8_4.word_wrap = True
    p8_4h = tf8_4.paragraphs[0]
    p8_4h.text = "4. AUTOMATED UPI PENNY-DROP & SPLIT REFUNDS"
    p8_4h.font.size = Pt(12)
    p8_4h.font.bold = True
    p8_4h.font.color.rgb = c_critical
    p8_4h.space_after = Pt(6)
    add_highlight_bullet(tf8_4, "• Frictionless Micro-Settlements:", "Algorithmic bank account verification via ₹1 NPCI UPI penny-drops, enabling automated fractional payouts and dynamic deductions for missing product accessories or tags.", 6)

    # =========================================================================
    # SLIDE 9: THANK YOU (LIGHT BACKGROUND TEMPLATE)
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    apply_slide_template(slide9)

    s9_num = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(1.3), Inches(0.36))
    s9_num.fill.solid()
    s9_num.fill.fore_color.rgb = RGBColor(241, 245, 249)
    s9_num.line.color.rgb = c_border
    tf_s9 = s9_num.text_frame
    p_s9 = tf_s9.paragraphs[0]
    p_s9.text = "SLIDE 9"
    p_s9.font.size = Pt(10.5)
    p_s9.font.bold = True
    p_s9.font.color.rgb = c_teal
    p_s9.alignment = PP_ALIGN.CENTER

    # Center Hero Card
    hero_card9 = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.5), Inches(10.9), Inches(5.0))
    hero_card9.fill.solid()
    hero_card9.fill.fore_color.rgb = c_card_white
    hero_card9.line.color.rgb = c_border

    tf9 = hero_card9.text_frame
    tf9.margin_left = Inches(0.5)
    tf9.margin_right = Inches(0.5)
    tf9.margin_top = Inches(0.8)
    tf9.word_wrap = True

    p_ty = tf9.paragraphs[0]
    p_ty.text = "THANK YOU"
    p_ty.font.size = Pt(56)
    p_ty.font.bold = True
    p_ty.font.color.rgb = c_text_dark
    p_ty.alignment = PP_ALIGN.CENTER
    p_ty.space_after = Pt(14)

    p_ty_sub = tf9.add_paragraph()
    p_ty_sub.text = "ReturnGuard: AI-Assisted Return-Fraud Detection Platform"
    p_ty_sub.font.size = Pt(20)
    p_ty_sub.font.color.rgb = c_text_muted
    p_ty_sub.alignment = PP_ALIGN.CENTER
    p_ty_sub.space_after = Pt(28)

    p_ty_open = tf9.add_paragraph()
    p_ty_open.text = "Open for Questions, Architectural Inquiries & Technical Discussion"
    p_ty_open.font.size = Pt(16)
    p_ty_open.font.bold = True
    p_ty_open.font.color.rgb = c_teal
    p_ty_open.alignment = PP_ALIGN.CENTER

    # Save logic
    output_pptx = "ReturnGuard_Presentation.pptx"
    try:
        prs.save(output_pptx)
        print(f"Successfully generated updated 9-slide PowerPoint presentation: {output_pptx}")
    except PermissionError:
        fallback_pptx = "ReturnGuard_Presentation_Updated.pptx"
        prs.save(fallback_pptx)
        print(f"Notice: '{output_pptx}' is currently open in PowerPoint. Saved updated deck to '{fallback_pptx}'.")

if __name__ == "__main__":
    create_presentation()
